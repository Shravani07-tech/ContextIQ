# tests/test_ingestion_phase_a.py
# Compact End-to-End Phase A ingestion test covering all 8 format categories.

import os
from unittest.mock import patch
import pytest

from ingest import chunk_documents, load_documents
from parsers import NormalizedDocument


def test_full_phase_a_ingestion_corpus(tmp_path):
    # 1. TXT
    (tmp_path / "sample.txt").write_text("ContextIQ 2.0 plain text content.", encoding="utf-8")

    # 2. PDF (mocked text extraction)
    (tmp_path / "sample.pdf").write_bytes(b"%PDF-1.4 dummy pdf header")

    # 3. DOCX
    import docx
    d = docx.Document()
    d.add_heading("Methodology", level=1)
    d.add_paragraph("ContextIQ Phase A ingestion pipeline testing.")
    d.save(str(tmp_path / "sample.docx"))

    # 4. PPTX
    from pptx import Presentation
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Phase A Presentation"
    prs.save(str(tmp_path / "sample.pptx"))

    # 5. XLSX
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "Score"])
    ws.append(["Accuracy", "99%"])
    wb.save(str(tmp_path / "sample.xlsx"))

    # 6. CSV
    (tmp_path / "sample.csv").write_text("ID,Value\n1,100\n2,200\n", encoding="utf-8")

    # 7. Markdown
    (tmp_path / "sample.md").write_text("# Title\n\nMarkdown document content.\n", encoding="utf-8")

    # 8. HTML
    (tmp_path / "sample.html").write_text("<html><body><h1>Header</h1><p>HTML document content.</p></body></html>", encoding="utf-8")

    # Mock PDFParser page text so pypdf doesn't return empty for dummy PDF bytes
    with patch("parsers.PDFParser.parse") as mock_pdf_parse:
        mock_pdf_parse.return_value = NormalizedDocument(
            filename="sample.pdf",
            text="Page 1 PDF text content for ContextIQ 2.0",
            file_type="pdf",
            document_id="sample.pdf",
            source_path=str(tmp_path / "sample.pdf"),
            pages=[(1, "Page 1 PDF text content for ContextIQ 2.0")],
            extraction_method="text",
        )

        # Run load_documents on tmp_path
        docs = load_documents(str(tmp_path))
        filenames = {d["filename"] for d in docs}

        expected_files = {
            "sample.txt",
            "sample.pdf",
            "sample.docx",
            "sample.pptx",
            "sample.xlsx",
            "sample.csv",
            "sample.md",
            "sample.html",
        }

        assert expected_files.issubset(filenames)

        # Run chunking
        chunks = chunk_documents(docs)
        assert len(chunks) > 0

        # Verify metadata fields are preserved
        chunk_files = {c["filename"] for c in chunks}
        assert expected_files.issubset(chunk_files)
