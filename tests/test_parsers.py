# tests/test_parsers.py
import os
import pytest
from parsers import (
    NormalizedDocument,
    ParserRegistry,
    TextParser,
    PDFParser,
    DOCXParser,
    PPTXParser,
    XLSXParser,
    default_registry,
)

def test_text_parser(tmp_path):
    txt_file = tmp_path / "sample.txt"
    txt_file.write_text("Hello ContextIQ 2.0!", encoding="utf-8")

    parser = TextParser()
    assert parser.can_parse("sample.txt") is True
    assert parser.can_parse("sample.pdf") is False

    doc = parser.parse(str(txt_file))
    assert isinstance(doc, NormalizedDocument)
    assert doc.filename == "sample.txt"
    assert doc.text == "Hello ContextIQ 2.0!"
    assert doc.file_type == "txt"
    assert doc.pages is None


def test_docx_parser(tmp_path):
    import docx

    docx_path = tmp_path / "sample.docx"
    doc_file = docx.Document()
    doc_file.add_heading("Architecture Overview", level=1)
    doc_file.add_paragraph("ContextIQ 2.0 supports modern document formats.")
    doc_file.save(str(docx_path))

    parser = DOCXParser()
    assert parser.can_parse("sample.docx") is True
    norm_doc = parser.parse(str(docx_path))

    assert norm_doc.file_type == "docx"
    assert norm_doc.filename == "sample.docx"
    assert "ContextIQ 2.0" in norm_doc.text
    assert norm_doc.sections is not None
    assert norm_doc.sections[0][0] == "Architecture Overview"


def test_pptx_parser(tmp_path):
    from pptx import Presentation

    pptx_path = tmp_path / "sample.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    title.text = "Phase A Presentation"
    prs.save(str(pptx_path))

    parser = PPTXParser()
    assert parser.can_parse("sample.pptx") is True
    norm_doc = parser.parse(str(pptx_path))

    assert norm_doc.file_type == "pptx"
    assert norm_doc.filename == "sample.pptx"
    assert "Phase A Presentation" in norm_doc.text
    assert norm_doc.slides is not None
    assert norm_doc.slides[0][0] == 1


def test_xlsx_parser(tmp_path):
    import openpyxl

    xlsx_path = tmp_path / "financials.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws.append(["Quarter", "Revenue", "Margin"])
    ws.append(["Q1 2026", "$1,200,000", "85%"])
    wb.save(str(xlsx_path))

    parser = XLSXParser()
    assert parser.can_parse("financials.xlsx") is True
    norm_doc = parser.parse(str(xlsx_path))

    assert norm_doc.file_type == "xlsx"
    assert norm_doc.filename == "financials.xlsx"
    assert "Revenue: $1,200,000" in norm_doc.text
    assert norm_doc.sheets is not None
    assert norm_doc.sheets[0][0] == "Revenue"


def test_parser_registry(tmp_path):
    registry = default_registry
    txt_file = tmp_path / "doc.txt"
    txt_file.write_text("Registry Test", encoding="utf-8")

    parser = registry.get_parser("doc.txt")
    assert parser is not None
    assert isinstance(parser, TextParser)

    doc = registry.parse_file(str(txt_file))
    assert doc.text == "Registry Test"

    with pytest.raises(ValueError, match="Unsupported file type"):
        registry.parse_file("unknown.xyz")
