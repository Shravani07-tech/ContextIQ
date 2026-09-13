# parsers.py
#
# Unified document parser architecture for ContextIQ 2.0.
# Translates various file formats into a single NormalizedDocument representation.

from abc import ABC, abstractmethod
from dataclasses import dataclass, field
import logging
import os
import re
from typing import Any, Dict, List, Optional, Tuple

from pypdf import PdfReader

logger = logging.getLogger(__name__)

# References / Bibliography heading detector for back-matter stripping.
_REFERENCES_HEADING = re.compile(
    r"(?im)^[ \t]*(?:\d+\.|[IVXLCDM]+\.)?[ \t]*(references|bibliography|works cited)(?:[ \t]+and[ \t]+(?:references|bibliography|works cited))?[ \t]*$"
)


def strip_references(text: str) -> str:
    """
    Remove a trailing References/Bibliography section if present in the latter half of the document.
    """
    for match in _REFERENCES_HEADING.finditer(text):
        if match.start() > len(text) * 0.5:
            return text[: match.start()]
    return text


@dataclass
class NormalizedDocument:
    """
    Normalized internal document model representing extracted content from any supported file type.
    """

    filename: str
    text: str
    file_type: str
    document_id: str
    source_path: Optional[str] = None
    pages: Optional[List[Tuple[int, str]]] = None  # [(page_num, text), ...]
    slides: Optional[List[Tuple[int, str]]] = None  # [(slide_num, text), ...]
    sheets: Optional[List[Tuple[str, str]]] = None  # [(sheet_name, text), ...]
    sections: Optional[List[Tuple[str, str]]] = None  # [(heading, text), ...]
    metadata: Dict[str, Any] = field(default_factory=dict)
    extraction_method: str = "text"  # "text" | "ocr"


class BaseParser(ABC):
    """Abstract base class for all file format parsers."""

    @abstractmethod
    def can_parse(self, filename: str) -> bool:
        """Return True if this parser supports the given filename/extension."""
        pass

    @abstractmethod
    def parse(self, file_path: str) -> NormalizedDocument:
        """Parse file at `file_path` and return a NormalizedDocument."""
        pass


class TextParser(BaseParser):
    """Parser for plain text files (.txt)."""

    def can_parse(self, filename: str) -> bool:
        return filename.lower().endswith(".txt")

    def parse(self, file_path: str) -> NormalizedDocument:
        filename = os.path.basename(file_path)
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()

        return NormalizedDocument(
            filename=filename,
            text=text,
            file_type="txt",
            document_id=filename,
            source_path=file_path,
            pages=None,
            extraction_method="text",
        )


class PDFParser(BaseParser):
    """Parser for PDF documents (.pdf)."""

    def can_parse(self, filename: str) -> bool:
        return filename.lower().endswith(".pdf")

    def parse(self, file_path: str) -> NormalizedDocument:
        from ocr_service import default_ocr_service

        filename = os.path.basename(file_path)
        reader = PdfReader(file_path)
        pages: List[Tuple[int, str]] = []

        for page_num, page in enumerate(reader.pages, 1):
            text = page.extract_text()
            if text and text.strip():
                pages.append((page_num, text))
            else:
                pages.append((page_num, ""))

        extraction_method = "text"
        non_empty_pages = [(p_num, txt) for p_num, txt in pages if txt.strip()]

        if not non_empty_pages:
            logger.info("No text found in '%s'; attempting local OCR fallback...", filename)
            ocr_pages, method = default_ocr_service.ocr_pdf(file_path)
            if ocr_pages:
                pages = ocr_pages
                extraction_method = method
            else:
                pages = non_empty_pages
        else:
            pages = non_empty_pages

        full_text = "\n".join(text for _, text in pages)

        return NormalizedDocument(
            filename=filename,
            text=full_text,
            file_type="pdf",
            document_id=filename,
            source_path=file_path,
            pages=pages if pages else None,
            extraction_method=extraction_method,
        )



class DOCXParser(BaseParser):
    """Parser for Microsoft Word documents (.docx)."""

    def can_parse(self, filename: str) -> bool:
        return filename.lower().endswith(".docx")

    def parse(self, file_path: str) -> NormalizedDocument:
        import docx

        filename = os.path.basename(file_path)
        doc = docx.Document(file_path)

        sections: List[Tuple[str, str]] = []
        current_heading = "Document Body"
        current_text_lines: List[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if p.style and p.style.name and p.style.name.startswith("Heading"):
                if current_text_lines:
                    sections.append((current_heading, "\n".join(current_text_lines)))
                    current_text_lines = []
                current_heading = text
            else:
                current_text_lines.append(text)

        for table in doc.tables:
            table_lines: List[str] = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_lines.append(" | ".join(row_cells))
            if table_lines:
                current_text_lines.append("\n".join(table_lines))

        if current_text_lines:
            sections.append((current_heading, "\n".join(current_text_lines)))

        full_text = "\n\n".join(f"=== {h} ===\n{t}" for h, t in sections) if sections else ""

        return NormalizedDocument(
            filename=filename,
            text=full_text,
            file_type="docx",
            document_id=filename,
            source_path=file_path,
            sections=sections if sections else None,
            extraction_method="text",
        )


class PPTXParser(BaseParser):
    """Parser for Microsoft PowerPoint documents (.pptx)."""

    def can_parse(self, filename: str) -> bool:
        return filename.lower().endswith(".pptx")

    def parse(self, file_path: str) -> NormalizedDocument:
        from pptx import Presentation

        filename = os.path.basename(file_path)
        prs = Presentation(file_path)
        slides: List[Tuple[int, str]] = []

        for slide_num, slide in enumerate(prs.slides, 1):
            slide_lines: List[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_lines.append(text)
                elif shape.has_table:
                    for row in shape.table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            slide_lines.append(" | ".join(row_cells))

            if slide_lines:
                slides.append((slide_num, "\n".join(slide_lines)))

        full_text = "\n\n".join(f"--- Slide {n} ---\n{txt}" for n, txt in slides)

        return NormalizedDocument(
            filename=filename,
            text=full_text,
            file_type="pptx",
            document_id=filename,
            source_path=file_path,
            slides=slides if slides else None,
            extraction_method="text",
        )


class XLSXParser(BaseParser):
    """Parser for Microsoft Excel workbooks (.xlsx)."""

    def can_parse(self, filename: str) -> bool:
        return filename.lower().endswith(".xlsx")

    def parse(self, file_path: str) -> NormalizedDocument:
        import openpyxl

        filename = os.path.basename(file_path)
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheets: List[Tuple[str, str]] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            header_idx = None
            headers = []
            for idx, r in enumerate(rows):
                non_empty = [str(cell).strip() for cell in r if cell is not None and str(cell).strip() != ""]
                if non_empty:
                    header_idx = idx
                    headers = [str(cell).strip() if cell is not None else f"Col_{i+1}" for i, cell in enumerate(r)]
                    break

            if header_idx is None:
                continue

            sheet_lines: List[str] = []
            for r in rows[header_idx + 1:]:
                row_fields = []
                for i, cell in enumerate(r):
                    if cell is not None and str(cell).strip() != "":
                        h_name = headers[i] if i < len(headers) and headers[i] else f"Col_{i+1}"
                        row_fields.append(f"{h_name}: {str(cell).strip()}")
                if row_fields:
                    sheet_lines.append(" | ".join(row_fields))

            if sheet_lines:
                sheets.append((sheet_name, "\n".join(sheet_lines)))

        wb.close()
        full_text = "\n\n".join(f"=== Sheet: {sname} ===\n{stxt}" for sname, stxt in sheets)

        return NormalizedDocument(
            filename=filename,
            text=full_text,
            file_type="xlsx",
            document_id=filename,
            source_path=file_path,
            sheets=sheets if sheets else None,
            extraction_method="text",
        )


class CSVParser(BaseParser):
    """Parser for CSV datasets (.csv)."""

    def can_parse(self, filename: str) -> bool:
        return filename.lower().endswith(".csv")

    def parse(self, file_path: str) -> NormalizedDocument:
        import csv

        filename = os.path.basename(file_path)
        content = ""
        for encoding in ["utf-8", "latin-1", "cp1252"]:
            try:
                with open(file_path, "r", encoding=encoding) as f:
                    content = f.read()
                break
            except Exception:
                continue

        lines = [line for line in content.splitlines() if line.strip()]
        if not lines:
            return NormalizedDocument(
                filename=filename, text="", file_type="csv", document_id=filename, source_path=file_path
            )

        sample = "\n".join(lines[:10])
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
            delimiter = dialect.delimiter
        except Exception:
            delimiter = ","

        reader = csv.reader(lines, delimiter=delimiter)
        rows = list(reader)
        if not rows:
            return NormalizedDocument(
                filename=filename, text="", file_type="csv", document_id=filename, source_path=file_path
            )

        headers = [h.strip() if h.strip() else f"Col_{i+1}" for i, h in enumerate(rows[0])]
        data_rows: List[str] = []

        for r_idx, row in enumerate(rows[1:], 1):
            fields = []
            for c_idx, val in enumerate(row):
                val_str = val.strip()
                if val_str:
                    h_name = headers[c_idx] if c_idx < len(headers) else f"Col_{c_idx+1}"
                    fields.append(f"{h_name}: {val_str}")
            if fields:
                data_rows.append(f"Row {r_idx}: " + " | ".join(fields))

        full_text = "\n".join(data_rows)

        return NormalizedDocument(
            filename=filename,
            text=full_text,
            file_type="csv",
            document_id=filename,
            source_path=file_path,
            extraction_method="text",
        )


class MarkdownParser(BaseParser):
    """Parser for Markdown documents (.md, .markdown)."""

    def can_parse(self, filename: str) -> bool:
        lower = filename.lower()
        return lower.endswith(".md") or lower.endswith(".markdown")

    def parse(self, file_path: str) -> NormalizedDocument:
        filename = os.path.basename(file_path)
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()

        sections: List[Tuple[str, str]] = []
        current_heading = "Overview"
        current_lines: List[str] = []

        for line in content.splitlines():
            stripped = line.strip()
            if stripped.startswith("#"):
                heading_text = stripped.lstrip("#").strip()
                if current_lines:
                    sections.append((current_heading, "\n".join(current_lines)))
                    current_lines = []
                if heading_text:
                    current_heading = heading_text
            else:
                current_lines.append(line)

        if current_lines:
            sections.append((current_heading, "\n".join(current_lines)))

        full_text = "\n\n".join(f"### Section: {h}\n{t}" for h, t in sections) if sections else content

        return NormalizedDocument(
            filename=filename,
            text=full_text,
            file_type="md",
            document_id=filename,
            source_path=file_path,
            sections=sections if sections else None,
            extraction_method="text",
        )


class HTMLParser(BaseParser):
    """Parser for HTML documents (.html, .htm)."""

    def can_parse(self, filename: str) -> bool:
        lower = filename.lower()
        return lower.endswith(".html") or lower.endswith(".htm")

    def parse(self, file_path: str) -> NormalizedDocument:
        from bs4 import BeautifulSoup

        filename = os.path.basename(file_path)
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            raw_html = f.read()

        soup = BeautifulSoup(raw_html, "html.parser")

        for tag in soup(["script", "style", "nav", "footer", "header", "noscript"]):
            tag.decompose()

        title_tag = soup.find("title")
        doc_title = title_tag.get_text().strip() if title_tag else filename

        sections: List[Tuple[str, str]] = []
        current_heading = doc_title
        current_lines: List[str] = []

        body = soup.find("body") or soup
        for element in body.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "table", "ul", "ol"]):
            tag_name = element.name.lower()
            if tag_name.startswith("h"):
                text = element.get_text().strip()
                if text:
                    if current_lines:
                        sections.append((current_heading, "\n".join(current_lines)))
                        current_lines = []
                    current_heading = text
            elif tag_name == "table":
                table_lines = []
                for tr in element.find_all("tr"):
                    cells = [td.get_text().strip() for td in tr.find_all(["td", "th"]) if td.get_text().strip()]
                    if cells:
                        table_lines.append(" | ".join(cells))
                if table_lines:
                    current_lines.append("\n".join(table_lines))
            else:
                text = element.get_text().strip()
                if text:
                    current_lines.append(text)

        if current_lines:
            sections.append((current_heading, "\n".join(current_lines)))

        full_text = "\n\n".join(f"=== {h} ===\n{t}" for h, t in sections) if sections else soup.get_text(separator="\n")

        return NormalizedDocument(
            filename=filename,
            text=full_text,
            file_type="html",
            document_id=filename,
            source_path=file_path,
            sections=sections if sections else None,
            extraction_method="text",
        )


class ParserRegistry:
    """Registry managing available document parsers."""

    def __init__(self) -> None:
        self._parsers: List[BaseParser] = [
            TextParser(),
            PDFParser(),
            DOCXParser(),
            PPTXParser(),
            XLSXParser(),
            CSVParser(),
            MarkdownParser(),
            HTMLParser(),
        ]

    def register_parser(self, parser: BaseParser) -> None:
        """Register a new parser instance."""
        self._parsers.append(parser)

    def get_parser(self, filename: str) -> Optional[BaseParser]:
        """Find the matching parser for a given filename."""
        for parser in self._parsers:
            if parser.can_parse(filename):
                return parser
        return None

    def parse_file(self, file_path: str) -> NormalizedDocument:
        """Parse a file using the registered parsers."""
        filename = os.path.basename(file_path)
        parser = self.get_parser(filename)
        if not parser:
            raise ValueError(f"Unsupported file type: {filename}")
        return parser.parse(file_path)


# Global registry singleton instance
default_registry = ParserRegistry()


